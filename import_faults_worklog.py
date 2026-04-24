#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Import worklog xlsx rows into fault_reports."""

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
from init_db import get_db_path
from import_review_support import (
    build_source_record_key,
    create_import_batch,
    enqueue_fault_review_item,
    ensure_station_name_proposal,
    fault_source_key_exists,
    get_columns,
    get_project_row,
    multi_project_import_enabled,
    project_code_from_system_type,
    resolve_station_match,
    split_station_tokens,
    update_import_batch_stats,
)
from utils import backup_sqlite_database, create_db_connection
from worklog_fault_types import classify_worklog_entry


SOURCE_FILE = r"e:\办公\图像监控\工作记录.xlsx"
SOURCE_TYPE = "import_worklog"
SOURCE_SYSTEM = "worklog"
IMPORT_MODE = "best-effort"
TIMEZONE_DEFAULT = "Asia/Shanghai"
SYSTEM_TYPE_ALIASES = {
    "智慧监控": "图像监控",
    "智慧巡视": "智能巡视",
}
TARGET_TYPES = {"图像监控", "智能巡视", "辅控系统", "智慧监控", "智慧巡视"}
SKIP_STATIONS = {"缙云", "遂昌公司", "丽水", "湖州", "钦矿变", "钼矿变"}


FAIL_ON_ANY_FAILURE = "any_failure"
FAIL_ON_DUPLICATE = "duplicate"
FAIL_ON_ROWS_SKIPPED = "rows_skipped"


class ImportAbortError(RuntimeError):
    def __init__(self, message, *, report=None):
        super().__init__(message)
        self.report = report


def parse_fail_on_rules(value):
    if not value:
        return set()
    rules = set()
    for part in str(value).split(","):
        normalized = part.strip().lower().replace("-", "_")
        if normalized:
            rules.add(normalized)
    return rules


def should_abort_import(*, fail_on_rules, reason, action):
    if not fail_on_rules:
        return False
    normalized_reason = str(reason).strip().lower().replace("-", "_")
    normalized_action = str(action).strip().lower().replace("-", "_")
    if FAIL_ON_ANY_FAILURE in fail_on_rules and normalized_action in {"queue", "skip"}:
        return True
    if normalized_reason in fail_on_rules or normalized_action in fail_on_rules:
        return True
    if normalized_action == "duplicate_skip" and FAIL_ON_DUPLICATE in fail_on_rules:
        return True
    if normalized_action == "skip" and FAIL_ON_ROWS_SKIPPED in fail_on_rules:
        return True
    return False


def ensure_worklog_fault_columns(conn):
    existing_columns = get_columns(conn, "fault_reports")
    required_columns = {
        "camera_location_text": "TEXT",
        "camera_slot_id": "INTEGER",
        "project_device_code": "TEXT",
    }
    for column_name, column_type in required_columns.items():
        if column_name in existing_columns:
            continue
        conn.execute(f"ALTER TABLE fault_reports ADD COLUMN {column_name} {column_type}")
        existing_columns.add(column_name)


def abort_import_if_needed(*, fail_on_rules, row_result, reason, action):
    if not should_abort_import(fail_on_rules=fail_on_rules, reason=reason, action=action):
        return
    row_index = row_result.get("row_index")
    raise ImportAbortError(
        f"fail-on rule triggered: {reason} at row {row_index}" if row_index else f"fail-on rule triggered: {reason}"
    )


def build_report(
    *,
    database_path,
    source_file,
    dry_run,
    fail_on_rules,
    timezone_default,
    stats,
    report_rows,
    aborted=False,
    error=None,
):
    report = {
        "database": str(database_path),
        "source": str(source_file),
        "source_type": SOURCE_TYPE,
        "mode": IMPORT_MODE,
        "dry_run": dry_run,
        "fail_on": sorted(fail_on_rules),
        "timezone_default": timezone_default,
        "aborted": aborted,
        **stats,
        "rows": report_rows,
    }
    if error:
        report["error"] = error
    return report


def write_report_file(report_path, report):
    if not report_path:
        return
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def parse_time(value):
    if value is None:
        return None
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")

    text = str(value).strip()
    if not text:
        return None
    text = text.split("-")[0].split("~")[0].split("至")[0].strip()
    match = re.match(r"(\d{4})年(\d{1,2})月(\d{1,2})日", text)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    return None


def resolve_fault_type_payload(content):
    fault_type = classify_worklog_entry(content)
    return {
        "is_fault": fault_type["is_fault"],
        "non_fault_reason": fault_type["reason"] if not fault_type["is_fault"] else None,
        "fault_type": fault_type["type_label"] if fault_type["is_fault"] else None,
        "fault_type_code": fault_type["type_code"] if fault_type["is_fault"] else None,
    }


def infer_fault_type(content):
    payload = resolve_fault_type_payload(content)
    return payload["fault_type"]


def normalize_worklog_system_type(value):
    text = str(value or "").strip()
    return SYSTEM_TYPE_ALIASES.get(text, text)


def make_legacy_idempotency_key(station_id, time_text, description):
    raw = f"{station_id}|{time_text or ''}|{description or ''}"
    return hashlib.md5(raw.encode("utf-8")).hexdigest()[:24]


def build_description(location, content):
    parts = []
    if content:
        parts.append(content)
    if location and location not in str(content or ""):
        parts.append(f"地点: {location}")
    return " | ".join(parts) if parts else (content or "")


def normalize_camera_location(value):
    text = str(value or "").strip().lower()
    for token in [
        "更换",
        "维修",
        "检修",
        "抢修",
        "调试",
        "处理",
        "恢复",
        "恢复正常",
        "后正常",
        "重新",
        "摄像机",
        "摄像头",
        "球机",
        "枪机",
        "监控点",
        "视频",
        "图像",
        "恢复正常",
        "故障",
        "异常",
        "接口",
        "集中电源",
        "地点",
        " ",
        "\t",
        "\r",
        "\n",
        "_",
        "-",
        "/",
        "\\",
        "|",
        "，",
        ",",
        "。",
        ".",
        "：",
        ":",
        "；",
        ";",
        "（",
        "）",
        "(",
        ")",
        "[",
        "]",
    ]:
        text = text.replace(token, "")
    return text


def _looks_like_camera_location(value):
    text = str(value or "").strip()
    if not text:
        return False
    keywords = (
        "侧",
        "室",
        "场地",
        "门",
        "通道",
        "开关",
        "主变",
        "电容器",
        "蓄电池",
        "楼",
        "围墙",
        "大门",
        "#",
        "kV",
        "kv",
    )
    return any(keyword in text for keyword in keywords)


def derive_camera_location_text(location, content):
    if _looks_like_camera_location(location):
        return str(location).strip()

    content_text = str(content or "").strip()
    if not content_text:
        return ""

    match = re.search(
        r"([\u4e00-\u9fffA-Za-z0-9#-]+?(?:侧|室|场地|门口|大门|通道))(?:\s*[-#]?\d+\s*[#号]?)?\s*(?:摄像机|摄像头|球机|枪机)",
        content_text,
    )
    if match:
        return match.group(1).strip()

    return ""


def resolve_camera_binding(conn, *, station_id, project_id, location, content):
    if not project_id:
        return None
    if "camera_slots" not in [row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()]:
        return None

    derived_location = derive_camera_location_text(location, content)
    normalized_hints = {
        normalize_camera_location(location),
        normalize_camera_location(derived_location),
        normalize_camera_location(content),
    }
    normalized_hints.discard("")
    if not normalized_hints:
        return None

    rows = conn.execute(
        """
        SELECT
            cs.id AS slot_id,
            cs.location_desc,
            cs.area,
            c.id AS camera_id,
            c.location_desc AS camera_location,
            c.area AS camera_area,
            c.project_camera_code,
            c.camera_index
        FROM camera_slots cs
        LEFT JOIN cameras c
          ON c.slot_id = cs.id
         AND c.status = 'active'
        WHERE cs.station_id = ?
          AND cs.project_id = ?
        ORDER BY cs.id
        """,
        (station_id, project_id),
    ).fetchall()

    matches = []
    for row in rows:
        best_candidate = ""
        best_length = 0
        for candidate in (
            row["camera_location"] if hasattr(row, "keys") else row[4],
            row["location_desc"] if hasattr(row, "keys") else row[1],
            row["camera_area"] if hasattr(row, "keys") else row[5],
            row["area"] if hasattr(row, "keys") else row[2],
        ):
            normalized_candidate = normalize_camera_location(candidate)
            if not normalized_candidate:
                continue
            if any(
                normalized_candidate in hint or hint in normalized_candidate
                for hint in normalized_hints
            ):
                if len(normalized_candidate) > best_length:
                    best_candidate = str(candidate or "").strip()
                    best_length = len(normalized_candidate)
        if best_length > 0:
            matches.append(
                {
                    "slot_id": row["slot_id"] if hasattr(row, "keys") else row[0],
                    "camera_id": row["camera_id"] if hasattr(row, "keys") else row[3],
                    "location_desc": best_candidate,
                    "project_device_code": (
                        row["project_camera_code"] if hasattr(row, "keys") else row[6]
                    ) or ((row["camera_index"] if hasattr(row, "keys") else row[7]) or None),
                    "score": best_length,
                }
            )

    if not matches:
        return None

    matches.sort(key=lambda item: (-item["score"], item["slot_id"]))
    top_score = matches[0]["score"]
    top_matches = [item for item in matches if item["score"] == top_score]
    unique_matches = {item["slot_id"]: item for item in top_matches}
    if len(unique_matches) != 1:
        return None
    return next(iter(unique_matches.values()))


def get_or_create_project_batch(
    conn,
    batch_cache,
    project_code,
    *,
    dry_run=False,
    timezone_default=TIMEZONE_DEFAULT,
    report_path=None,
):
    if project_code in batch_cache:
        return batch_cache[project_code]

    project = get_project_row(conn, project_code)
    if not project:
        return None

    batch_id = None
    if not dry_run:
        batch_id = create_import_batch(
            conn,
            project_id=project["id"],
            source_type=SOURCE_TYPE,
            mode=IMPORT_MODE,
            file_count=1,
            report_path=str(report_path) if report_path else None,
            timezone_default_used=timezone_default,
        )
    batch_cache[project_code] = {
        "project": project,
        "batch_id": batch_id,
        "success_count": 0,
        "fail_count": 0,
    }
    return batch_cache[project_code]


def insert_fault_report(
    conn,
    *,
    station_id,
    system_type,
    fault_type,
    fault_type_code=None,
    description,
    handler_name,
    created_at,
    legacy_idempotency_key,
    project_row=None,
    batch_id=None,
    source_record_key=None,
    camera_id=None,
    camera_slot_id=None,
    camera_location_text=None,
    project_device_code=None,
    raw_time_value=None,
    timezone_default=TIMEZONE_DEFAULT,
):
    columns = get_columns(conn, "fault_reports")
    insert_columns = [
        "station_id",
        "system_type",
        "fault_type",
        "description",
        "reporter_name",
        "handler_name",
        "status",
        "closed_at",
        "created_at",
        "updated_at",
        "idempotency_key",
    ]
    values = [
        station_id,
        system_type,
        fault_type,
        description,
        "工作记录导入",
        handler_name,
        "closed",
        created_at,
        created_at,
        created_at,
        legacy_idempotency_key,
    ]

    optional_values = [
        ("camera_id", camera_id),
        ("project_id", project_row["id"] if project_row else None),
        ("camera_slot_id", camera_slot_id),
        ("source_type", SOURCE_TYPE),
        ("source_batch_id", str(batch_id) if batch_id else None),
        ("source_record_key", source_record_key),
        ("project_device_code", project_device_code),
        ("camera_location_text", camera_location_text),
        ("fault_type_label_snapshot", fault_type),
        ("fault_type_code", fault_type_code),
        ("fault_type_version_id", project_row.get("fault_type_version_id") if project_row else None),
        ("source_time_raw", str(raw_time_value) if raw_time_value is not None else None),
        ("source_timezone", timezone_default),
    ]
    for column_name, column_value in optional_values:
        if column_name in columns:
            insert_columns.append(column_name)
            values.append(column_value)

    placeholders = ", ".join(["?"] * len(insert_columns))
    conn.execute(
        f"""
        INSERT INTO fault_reports ({", ".join(insert_columns)})
        VALUES ({placeholders})
        """,
        values,
    )


def import_worklog_file(
    source_file=SOURCE_FILE,
    *,
    database_path=None,
    dry_run=False,
    report_path=None,
    fail_on="",
    timezone_default=TIMEZONE_DEFAULT,
):
    workbook = openpyxl.load_workbook(source_file, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    database_path = Path(database_path or get_db_path()).resolve()
    source_path = Path(source_file).resolve()
    conn = create_db_connection(database_path, row_factory=True, enable_wal=True)
    ensure_worklog_fault_columns(conn)
    multi_project_enabled = multi_project_import_enabled(conn)
    batch_cache = {}
    fail_on_rules = parse_fail_on_rules(fail_on)
    report_rows = []
    stats = {
        "inserted": 0,
        "duplicates_skipped": 0,
        "queue_items_created": 0,
        "station_proposals_created": 0,
        "rows_skipped": 0,
        "non_fault_rows_skipped": 0,
        "fail_count": 0,
        "errors": [],
    }

    try:
        for row_index, row in enumerate(rows[2:], start=3):
            if not row or row[5] is None:
                continue
            system_type = normalize_worklog_system_type(row[5])
            if system_type not in TARGET_TYPES:
                continue

            seq = row[0]
            raw_time = row[1]
            created_at = parse_time(raw_time)
            station_text = str(row[2]).strip() if row[2] else ""
            location = str(row[3]).strip() if row[3] else ""
            content = str(row[4]).strip() if row[4] else ""
            handler_name = str(row[7]).strip() if len(row) > 7 and row[7] else None
            fault_type_payload = resolve_fault_type_payload(content)
            if not fault_type_payload["is_fault"]:
                row_result = {
                    "row_index": row_index,
                    "sequence": seq,
                    "action": "skip",
                    "reason": "non_fault_work_item",
                    "non_fault_reason": fault_type_payload["non_fault_reason"],
                }
                stats["rows_skipped"] += 1
                stats["non_fault_rows_skipped"] += 1
                report_rows.append(row_result)
                abort_import_if_needed(
                    fail_on_rules=fail_on_rules,
                    row_result=row_result,
                    reason="non_fault_work_item",
                    action="skip",
                )
                continue
            fault_type = fault_type_payload["fault_type"]
            fault_type_code = fault_type_payload["fault_type_code"]
            description = build_description(location, content)

            if not station_text:
                row_result = {"row_index": row_index, "sequence": seq, "action": "skip", "reason": "missing_station_name"}
                stats["rows_skipped"] += 1
                report_rows.append(row_result)
                abort_import_if_needed(
                    fail_on_rules=fail_on_rules,
                    row_result=row_result,
                    reason="missing_station_name",
                    action="skip",
                )
                continue

            project_code = project_code_from_system_type(system_type)
            batch_state = (
                get_or_create_project_batch(
                    conn,
                    batch_cache,
                    project_code,
                    dry_run=dry_run,
                    timezone_default=timezone_default,
                    report_path=Path(report_path).resolve() if report_path else None,
                )
                if multi_project_enabled and project_code
                else None
            )
            batch_id = batch_state["batch_id"] if batch_state else None

            for station_token in split_station_tokens(station_text):
                row_result = {"row_index": row_index, "sequence": seq, "station_token": station_token}

                if station_token in SKIP_STATIONS:
                    stats["rows_skipped"] += 1
                    row_result["action"] = "skip"
                    row_result["reason"] = "station_skipped"
                    report_rows.append(row_result)
                    abort_import_if_needed(
                        fail_on_rules=fail_on_rules,
                        row_result=row_result,
                        reason="station_skipped",
                        action="skip",
                    )
                    continue

                source_record_key = (
                    build_source_record_key(
                        project_code or "legacy",
                        SOURCE_TYPE,
                        raw_external_id=f"{SOURCE_SYSTEM}:{seq}:{station_token}",
                    )
                    if multi_project_enabled
                    else None
                )
                match = resolve_station_match(
                    conn,
                    station_token,
                    source_system=SOURCE_SYSTEM,
                )

                raw_context = {
                    "sequence": seq,
                    "system_type": system_type,
                    "station_text": station_text,
                    "station_token": station_token,
                    "location": location,
                    "content": content,
                    "description": description,
                    "fault_type": fault_type,
                    "fault_type_code": fault_type_code,
                    "handler_name": handler_name,
                    "parsed_time": created_at,
                    "raw_time": str(raw_time) if raw_time is not None else None,
                    "source_timezone": timezone_default,
                }

                if match["should_create_proposal"] and multi_project_enabled and batch_state and not dry_run:
                    proposal_id = ensure_station_name_proposal(
                        conn,
                        import_batch_id=batch_id,
                        project_id=batch_state["project"]["id"],
                        source_system=SOURCE_SYSTEM,
                        external_name=station_token,
                        candidate_station_id=match["proposal_candidate_station_id"],
                        confidence_score=match["confidence_score"],
                        raw_context=raw_context,
                    )
                    if proposal_id:
                        stats["station_proposals_created"] += 1

                if not match["matched"]:
                    stats["fail_count"] += 1
                    row_result["action"] = "queue"
                    row_result["reason"] = "station_not_resolved"
                    report_rows.append(row_result)
                    if multi_project_enabled and batch_state and not dry_run:
                        queue_id = enqueue_fault_review_item(
                            conn,
                            import_batch_id=batch_id,
                            project_id=batch_state["project"]["id"],
                            source_type=SOURCE_TYPE,
                            source_record_key_candidate=source_record_key,
                            raw_payload=raw_context,
                            issue_type="station_not_resolved",
                            issue_detail=f"Unable to resolve station token: {station_token}",
                        )
                        if queue_id:
                            stats["queue_items_created"] += 1
                        batch_state["fail_count"] += 1
                    else:
                        stats["errors"].append(
                            {"sequence": seq, "station_token": station_token, "error": "station_not_resolved"}
                        )
                    abort_import_if_needed(
                        fail_on_rules=fail_on_rules,
                        row_result=row_result,
                        reason="station_not_resolved",
                        action="queue",
                    )
                    continue

                if multi_project_enabled and source_record_key and fault_source_key_exists(conn, source_record_key):
                    stats["duplicates_skipped"] += 1
                    row_result["action"] = "duplicate-skip"
                    row_result["reason"] = "source_record_key_exists"
                    report_rows.append(row_result)
                    abort_import_if_needed(
                        fail_on_rules=fail_on_rules,
                        row_result=row_result,
                        reason="source_record_key_exists",
                        action="duplicate_skip",
                    )
                    continue

                legacy_idempotency_key = make_legacy_idempotency_key(
                    match["station_id"], created_at, description
                )
                existing = conn.execute(
                    "SELECT id FROM fault_reports WHERE idempotency_key = ?",
                    (legacy_idempotency_key,),
                ).fetchone()
                if existing:
                    stats["duplicates_skipped"] += 1
                    row_result["action"] = "duplicate-skip"
                    row_result["reason"] = "legacy_idempotency_key_exists"
                    report_rows.append(row_result)
                    abort_import_if_needed(
                        fail_on_rules=fail_on_rules,
                        row_result=row_result,
                        reason="legacy_idempotency_key_exists",
                        action="duplicate_skip",
                    )
                    continue

                camera_binding = None
                camera_location_text = derive_camera_location_text(location, content)
                if batch_state:
                    camera_binding = resolve_camera_binding(
                        conn,
                        station_id=match["station_id"],
                        project_id=batch_state["project"]["id"],
                        location=location,
                        content=content,
                    )
                    if camera_binding and camera_binding.get("location_desc"):
                        camera_location_text = camera_binding["location_desc"]

                try:
                    if not dry_run:
                        insert_fault_report(
                            conn,
                            station_id=match["station_id"],
                            system_type=system_type,
                            fault_type=fault_type,
                            fault_type_code=fault_type_code,
                            description=description,
                            handler_name=handler_name,
                            created_at=created_at,
                            legacy_idempotency_key=legacy_idempotency_key,
                            project_row=batch_state["project"] if batch_state else None,
                            batch_id=batch_id,
                            source_record_key=source_record_key,
                            camera_id=camera_binding["camera_id"] if camera_binding else None,
                            camera_slot_id=camera_binding["slot_id"] if camera_binding else None,
                            camera_location_text=camera_location_text or None,
                            project_device_code=camera_binding["project_device_code"] if camera_binding else None,
                            raw_time_value=raw_time,
                            timezone_default=timezone_default,
                        )
                    stats["inserted"] += 1
                    row_result["action"] = "insert"
                    row_result["station_id"] = match["station_id"]
                    if camera_binding:
                        row_result["camera_slot_id"] = camera_binding["slot_id"]
                        row_result["camera_id"] = camera_binding["camera_id"]
                    if camera_location_text:
                        row_result["camera_location_text"] = camera_location_text
                    report_rows.append(row_result)
                    if batch_state:
                        batch_state["success_count"] += 1
                except Exception as exc:
                    stats["fail_count"] += 1
                    row_result["action"] = "queue"
                    row_result["reason"] = "insert_failed"
                    row_result["error"] = str(exc)
                    report_rows.append(row_result)
                    stats["errors"].append(
                        {
                            "sequence": seq,
                            "station_token": station_token,
                            "error": str(exc),
                        }
                    )
                    if multi_project_enabled and batch_state and not dry_run:
                        queue_id = enqueue_fault_review_item(
                            conn,
                            import_batch_id=batch_id,
                            project_id=batch_state["project"]["id"],
                            source_type=SOURCE_TYPE,
                            source_record_key_candidate=source_record_key,
                            raw_payload=raw_context,
                            issue_type="insert_failed",
                            issue_detail=str(exc),
                        )
                        if queue_id:
                            stats["queue_items_created"] += 1
                        batch_state["fail_count"] += 1
                    abort_import_if_needed(
                        fail_on_rules=fail_on_rules,
                        row_result=row_result,
                        reason="insert_failed",
                        action="queue",
                    )

        if not dry_run:
            for state in batch_cache.values():
                if state["batch_id"] is None:
                    continue
                update_import_batch_stats(
                    conn,
                    state["batch_id"],
                    success_count=state["success_count"],
                    fail_count=state["fail_count"],
                )
            conn.commit()
    except Exception as exc:
        if not dry_run:
            conn.rollback()
        report = build_report(
            database_path=database_path,
            source_file=source_path,
            dry_run=dry_run,
            fail_on_rules=fail_on_rules,
            timezone_default=timezone_default,
            stats=stats,
            report_rows=report_rows,
            aborted=True,
            error=str(exc),
        )
        write_report_file(Path(report_path).resolve() if report_path else None, report)
        if isinstance(exc, ImportAbortError):
            exc.report = report
        raise
    finally:
        conn.close()

    report = build_report(
        database_path=database_path,
        source_file=source_path,
        dry_run=dry_run,
        fail_on_rules=fail_on_rules,
        timezone_default=timezone_default,
        stats=stats,
        report_rows=report_rows,
    )
    write_report_file(Path(report_path).resolve() if report_path else None, report)
    return report


def backfill_worklog_camera_bindings(*, database_path=None, dry_run=False):
    database_path = Path(database_path or get_db_path()).resolve()
    conn = create_db_connection(database_path, row_factory=True, enable_wal=True)
    ensure_worklog_fault_columns(conn)
    stats = {
        "scanned": 0,
        "updated": 0,
        "slot_bound": 0,
        "location_filled": 0,
    }

    try:
        rows = conn.execute(
            """
            SELECT id, station_id, project_id, description, handler_note, camera_id, camera_slot_id, camera_location_text
            FROM fault_reports
            WHERE source_type = ?
              AND (
                    camera_id IS NULL
                 OR camera_slot_id IS NULL
                 OR camera_location_text IS NULL
                 OR TRIM(camera_location_text) = ''
              )
            ORDER BY id
            """,
            (SOURCE_TYPE,),
        ).fetchall()

        for row in rows:
            stats["scanned"] += 1
            description = str(row["description"] or "")
            if not description.strip():
                description = str(row["handler_note"] or "")
            content, _, suffix = description.partition(" | 地点:")
            location = suffix.strip()
            content = content.strip()

            camera_binding = resolve_camera_binding(
                conn,
                station_id=row["station_id"],
                project_id=row["project_id"],
                location=location,
                content=content,
            ) if row["project_id"] else None
            camera_location_text = derive_camera_location_text(location, content)
            if camera_binding and camera_binding.get("location_desc"):
                camera_location_text = camera_binding["location_desc"]

            update_fields = []
            update_values = []

            if camera_binding and row["camera_id"] != camera_binding["camera_id"]:
                update_fields.append("camera_id = ?")
                update_values.append(camera_binding["camera_id"])
            if camera_binding and row["camera_slot_id"] != camera_binding["slot_id"]:
                update_fields.append("camera_slot_id = ?")
                update_values.append(camera_binding["slot_id"])
            if camera_binding and camera_binding.get("project_device_code"):
                update_fields.append("project_device_code = COALESCE(project_device_code, ?)")
                update_values.append(camera_binding["project_device_code"])
            if camera_location_text and str(row["camera_location_text"] or "").strip() != camera_location_text:
                update_fields.append("camera_location_text = ?")
                update_values.append(camera_location_text)

            if not update_fields:
                continue

            if camera_binding:
                stats["slot_bound"] += 1
            if camera_location_text:
                stats["location_filled"] += 1
            stats["updated"] += 1

            if dry_run:
                continue

            update_values.append(row["id"])
            conn.execute(
                f"UPDATE fault_reports SET {', '.join(update_fields)} WHERE id = ?",
                update_values,
            )

        if dry_run:
            conn.rollback()
        else:
            conn.commit()
    finally:
        conn.close()

    return stats


def legacy_main_unused():
    backup_path = backup_sqlite_database(get_db_path(), label="import_worklog")
    if backup_path:
        print(f"[备份] 已创建数据库备份: {backup_path}")
    stats = import_worklog_file(SOURCE_FILE)
    print("=" * 60)
    print("Worklog import completed")
    print("=" * 60)
    for key, value in stats.items():
        if key == "errors":
            print(f"{key}: {len(value)}")
        else:
            print(f"{key}: {value}")
    return stats


def parse_args():
    parser = argparse.ArgumentParser(description="Import worklog xlsx rows into fault_reports")
    parser.add_argument("--source", default=SOURCE_FILE, help="Path to worklog xlsx file")
    parser.add_argument("--database", default=get_db_path(), help="Path to SQLite database")
    parser.add_argument("--report", help="Optional JSON report output path")
    parser.add_argument("--dry-run", action="store_true", help="Preview only, do not write database")
    parser.add_argument("--fail-on", default="", help="Comma-separated fail-fast rules")
    parser.add_argument("--timezone-default", default=TIMEZONE_DEFAULT, help="Timezone label stored on imported rows")
    parser.add_argument("--skip-backup", action="store_true", help="Skip automatic SQLite backup before write import")
    return parser.parse_args()


def main():
    args = parse_args()
    database_path = Path(args.database).resolve()
    report_path = Path(args.report).resolve() if args.report else None

    if not args.dry_run and not args.skip_backup:
        backup_path = backup_sqlite_database(database_path, label="import_worklog")
        if backup_path:
            print(f"[澶囦唤] 宸插垱寤烘暟鎹簱澶囦唤: {backup_path}")

    try:
        report = import_worklog_file(
            args.source,
            database_path=database_path,
            dry_run=args.dry_run,
            report_path=report_path,
            fail_on=args.fail_on,
            timezone_default=args.timezone_default,
        )
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return report
    except ImportAbortError as exc:
        if exc.report:
            print(json.dumps(exc.report, ensure_ascii=False, indent=2))
        else:
            print(json.dumps({"error": str(exc)}, ensure_ascii=False, indent=2))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
