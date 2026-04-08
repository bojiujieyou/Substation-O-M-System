#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Import daily fault summary Excel rows into fault_reports."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl
import xlrd

from ai_fault_analysis import (
    DailyFaultSummaryAIService,
    build_ai_trace,
    ensure_ai_runtime_schema,
    normalize_camera_hint,
)
from init_db import get_db_path
from import_review_support import (
    PROJECT_CODE_BY_SYSTEM_TYPE,
    build_source_record_key,
    create_import_batch,
    enqueue_fault_review_item,
    ensure_station_name_proposal,
    fault_source_key_exists,
    get_columns,
    get_project_row,
    resolve_station_match,
    table_exists,
    update_import_batch_stats,
)
from utils import create_db_connection


SOURCE_TYPE = "import_daily_fault_summary"
SOURCE_SYSTEM = "daily_fault_summary"
IMPORT_MODE = "best-effort"
TIMEZONE_DEFAULT = "Asia/Shanghai"


class DailyFaultSummaryParseError(ValueError):
    """Raised when the daily summary file cannot be parsed."""


def _reverse_project_type_mapping():
    return {project_code: system_type for system_type, project_code in PROJECT_CODE_BY_SYSTEM_TYPE.items()}


SYSTEM_TYPE_BY_PROJECT_CODE = _reverse_project_type_mapping()


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _load_excel_rows(source_file):
    path = Path(source_file)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        workbook = xlrd.open_workbook(str(path))
        sheet = workbook.sheet_by_index(0)
        return [sheet.row_values(index) for index in range(sheet.nrows)]
    if suffix == ".xlsx":
        workbook = openpyxl.load_workbook(path, data_only=True)
        try:
            worksheet = workbook.active
            return [list(row) for row in worksheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    raise DailyFaultSummaryParseError("只支持 .xls 或 .xlsx 文件")


def _parse_filename_date(source_path: Path):
    match = re.search(r"(20\d{2})(\d{2})(\d{2})", source_path.stem)
    if not match:
        return None
    year, month, day = map(int, match.groups())
    return datetime(year, month, day)


def _parse_source_date(rows, source_path: Path):
    filename_date = _parse_filename_date(source_path)
    default_year = filename_date.year if filename_date else datetime.now().year
    for row in rows[:5]:
        for cell in row[:3]:
            text = _cell_text(cell)
            if not text or "时间" not in text:
                continue
            match = re.search(r"时间[:：]?\s*(?:(\d{4})\s*年)?\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日", text)
            if not match:
                continue
            year = int(match.group(1) or default_year)
            month = int(match.group(2))
            day = int(match.group(3))
            return datetime(year, month, day), text
    if filename_date:
        return filename_date, filename_date.strftime("%Y-%m-%d")
    raise DailyFaultSummaryParseError("未识别到日报时间，请检查文件标题或文件名日期")


def _parse_reporter(rows):
    for row in rows[:5]:
        for cell in row:
            text = _cell_text(cell)
            match = re.search(r"检查人员[:：]\s*(.+)", text)
            if match:
                return match.group(1).strip()
    return "每日故障汇总导入"


def _find_header_row(rows):
    for index, row in enumerate(rows):
        first = _cell_text(row[0] if len(row) > 0 else "")
        second = _cell_text(row[1] if len(row) > 1 else "")
        if first == "变电站" and "问题描述" in second:
            return index
    raise DailyFaultSummaryParseError("未识别到“变电站 / 问题描述”表头")


def parse_daily_fault_summary(source_file):
    source_path = Path(source_file).resolve()
    rows = _load_excel_rows(source_path)
    if not rows:
        raise DailyFaultSummaryParseError("Excel 文件为空")

    source_date, source_time_raw = _parse_source_date(rows, source_path)
    reporter_name = _parse_reporter(rows)
    header_index = _find_header_row(rows)
    current_section = ""
    entries = []

    for row_index, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        station_name = _cell_text(row[0] if len(row) > 0 else "")
        problem_description = _cell_text(row[1] if len(row) > 1 else "")
        if not station_name and not problem_description:
            continue
        if station_name and not problem_description:
            current_section = station_name
            continue
        if not station_name:
            continue

        entries.append(
            {
                "row_index": row_index,
                "section": current_section,
                "station_name": station_name,
                "problem_description": problem_description,
            }
        )

    if not entries:
        raise DailyFaultSummaryParseError("未识别到有效的故障记录行")

    return {
        "source_path": source_path,
        "title": _cell_text(rows[0][0] if rows and rows[0] else source_path.stem),
        "source_date": source_date.strftime("%Y-%m-%d"),
        "source_time_raw": source_time_raw,
        "reporter_name": reporter_name,
        "entries": entries,
    }


def infer_fault_type(section, description):
    text = f"{section or ''} {description or ''}"
    if any(token in text for token in ["离线", "掉线", "断线"]):
        return "摄像头离线"
    if any(token in text for token in ["网络", "通信", "线路故障", "链路"]):
        return "网络故障"
    return "视频监控异常"


def build_fault_description(section, description):
    section_text = _cell_text(section)
    detail_text = _cell_text(description)
    if section_text and section_text not in detail_text:
        return f"{section_text} | {detail_text}"
    return detail_text


def normalize_camera_location(value):
    text = _cell_text(value).lower()
    for token in [
        "摄像头",
        "摄像机",
        "球机",
        "枪机",
        "枪",
        "球",
        "监控",
        "平台离线",
        "省公司",
        "（",
        "）",
        "(",
        ")",
        " ",
        "\t",
        "，",
        ",",
        "；",
        ";",
        "|",
    ]:
        text = text.replace(token, "")
    return text


def _normalize_text_for_semantic_dedupe(value):
    text = _cell_text(value).lower()
    for token in [
        " ",
        "\t",
        "\r",
        "\n",
        "_",
        "-",
        "/",
        "\\",
        "|",
        "，",
        ",",
        "。",
        ".",
        "：",
        ":",
        "；",
        ";",
        "（",
        "）",
        "(",
        ")",
        "【",
        "】",
        "[",
        "]",
    ]:
        text = text.replace(token, "")
    return text


def _normalize_fault_type_for_semantic_dedupe(value):
    text = _cell_text(value)
    if any(token in text for token in ["离线", "掉线", "断线"]):
        return "camera_offline"
    if any(token in text for token in ["网络", "通信", "链路"]):
        return "network_fault"
    if any(token in text for token in ["黑屏", "模糊", "视频", "图像", "画面"]):
        return "video_fault"
    return _normalize_text_for_semantic_dedupe(text)


def _build_semantic_camera_key(camera_location_text, description):
    candidate = normalize_camera_hint(camera_location_text) or normalize_camera_hint(description)
    return _normalize_text_for_semantic_dedupe(candidate)


def find_semantic_duplicate_fault(
    conn,
    *,
    station_id,
    project_id,
    source_date,
    fault_type,
    description,
    camera_location_text,
):
    semantic_camera_key = _build_semantic_camera_key(camera_location_text, description)
    if not semantic_camera_key:
        return None

    semantic_fault_type = _normalize_fault_type_for_semantic_dedupe(fault_type or description)
    columns = get_columns(conn, "fault_reports")
    if not columns:
        return None

    select_fields = [
        "id",
        "fault_type",
        "description",
        "created_at",
    ]
    if "fault_type_label_snapshot" in columns:
        select_fields.append("fault_type_label_snapshot")
    if "camera_location_text" in columns:
        select_fields.append("camera_location_text")

    query = f"""
        SELECT {", ".join(select_fields)}
        FROM fault_reports
        WHERE station_id = ?
          AND date(created_at) = date(?)
    """
    params = [station_id, source_date]

    if "project_id" in columns and project_id is not None:
        query += " AND project_id = ?"
        params.append(project_id)
    if "source_type" in columns:
        query += " AND source_type = ?"
        params.append(SOURCE_TYPE)

    rows = conn.execute(query, params).fetchall()
    for row in rows:
        row_fault_type = row["fault_type"] if hasattr(row, "keys") else row[1]
        if "fault_type_label_snapshot" in columns:
            snapshot = row["fault_type_label_snapshot"] if hasattr(row, "keys") else row[4]
            row_fault_type = snapshot or row_fault_type
        row_description = row["description"] if hasattr(row, "keys") else row[2]
        row_camera_location = ""
        if "camera_location_text" in columns:
            row_camera_location = row["camera_location_text"] if hasattr(row, "keys") else row[-1]

        if _normalize_fault_type_for_semantic_dedupe(row_fault_type) != semantic_fault_type:
            continue
        if _build_semantic_camera_key(row_camera_location, row_description) != semantic_camera_key:
            continue

        return row["id"] if hasattr(row, "keys") else row[0]

    return None


def resolve_camera_binding(conn, *, station_id, project_id, description):
    normalized_description = normalize_camera_location(description)
    if not normalized_description or not table_exists(conn, "camera_slots"):
        return None

    rows = conn.execute(
        """
        SELECT
            cs.id AS slot_id,
            cs.location_desc,
            cs.area,
            c.id AS camera_id,
            c.location_desc AS camera_location,
            c.area AS camera_area
        FROM camera_slots cs
        LEFT JOIN cameras c
          ON c.slot_id = cs.id
         AND c.status = 'active'
        WHERE cs.station_id = ?
          AND cs.project_id = ?
        ORDER BY cs.id
        """,
        (station_id, project_id),
    ).fetchall()

    matches = []
    for row in rows:
        candidates = [
            row["camera_location"] if hasattr(row, "keys") else row[4],
            row["location_desc"] if hasattr(row, "keys") else row[1],
            row["camera_area"] if hasattr(row, "keys") else row[5],
            row["area"] if hasattr(row, "keys") else row[2],
        ]
        normalized_candidates = {normalize_camera_location(value) for value in candidates if _cell_text(value)}
        normalized_candidates.discard("")
        if not normalized_candidates:
            continue
        if any(
            candidate in normalized_description or normalized_description in candidate
            for candidate in normalized_candidates
        ):
            matches.append(
                {
                    "slot_id": row["slot_id"] if hasattr(row, "keys") else row[0],
                    "camera_id": row["camera_id"] if hasattr(row, "keys") else row[3],
                    "location_desc": row["camera_location"] if hasattr(row, "keys") else row[4],
                }
            )

    unique_matches = {match["slot_id"]: match for match in matches}
    if len(unique_matches) != 1:
        return None
    return next(iter(unique_matches.values()))


def _derive_camera_location_text(problem_description, ai_result):
    source_text = _cell_text(problem_description)
    ai_text = ""
    if ai_result:
        ai_text = normalize_camera_hint(ai_result.get("camera_location_text"))
        if not ai_text:
            reason_hint = str(ai_result.get("reason") or "").strip()
            # Some models place the actual location into the reason field.
            if len(reason_hint) <= 40 and reason_hint and reason_hint in source_text:
                ai_text = normalize_camera_hint(reason_hint)
    if ai_text and ai_text in source_text:
        return ai_text
    return normalize_camera_hint(problem_description)


def build_idempotency_key(project_code, station_id, source_date, description):
    seed = f"{project_code}|{station_id}|{source_date}|{description}"
    return hashlib.md5(seed.encode("utf-8")).hexdigest()[:24]


def _system_type_for_project(project_code):
    return SYSTEM_TYPE_BY_PROJECT_CODE.get(project_code, "图像监控")


def _supports_daily_fault_import(conn):
    required_tables = {"projects", "import_batches", "fault_reports"}
    if not all(table_exists(conn, table_name) for table_name in required_tables):
        return False
    fault_columns = get_columns(conn, "fault_reports")
    required_columns = {"project_id", "source_type", "source_batch_id", "source_record_key", "fault_type_label_snapshot"}
    return required_columns.issubset(fault_columns)


def insert_fault_report(
    conn,
    *,
    station_id,
    project_id,
    batch_id,
    project_code,
    source_date,
    source_time_raw,
    reporter_name,
    description,
    fault_type,
    source_record_key,
    camera_id=None,
    camera_slot_id=None,
    camera_location_text=None,
    ai_confidence=None,
    ai_trace_json=None,
    timezone_default=TIMEZONE_DEFAULT,
):
    columns = get_columns(conn, "fault_reports")
    insert_columns = [
        "station_id",
        "system_type",
        "fault_type",
        "description",
        "reporter_name",
        "status",
        "created_at",
        "updated_at",
        "idempotency_key",
    ]
    values = [
        station_id,
        _system_type_for_project(project_code),
        fault_type,
        description,
        reporter_name,
        "open",
        source_date,
        source_date,
        build_idempotency_key(project_code, station_id, source_date, description),
    ]

    optional_values = [
        ("camera_id", camera_id),
        ("project_id", project_id),
        ("source_type", SOURCE_TYPE),
        ("source_batch_id", str(batch_id) if batch_id is not None else None),
        ("source_record_key", source_record_key),
        ("camera_slot_id", camera_slot_id),
        ("camera_location_text", camera_location_text),
        ("fault_type_label_snapshot", fault_type),
        ("source_time_raw", source_time_raw),
        ("source_timezone", timezone_default),
        ("ai_confidence", ai_confidence),
        ("ai_trace_json", ai_trace_json),
    ]
    for column_name, column_value in optional_values:
        if column_name in columns:
            insert_columns.append(column_name)
            values.append(column_value)

    placeholders = ", ".join(["?"] * len(insert_columns))
    conn.execute(
        f"""
        INSERT INTO fault_reports ({", ".join(insert_columns)})
        VALUES ({placeholders})
        """,
        values,
    )


def _build_file_result(source_name, summary):
    inserted = int(summary.get("inserted") or 0)
    failed = int(summary.get("fail_count") or 0)
    queued = int(summary.get("queue_items_created") or 0)
    duplicates = int(summary.get("duplicates_skipped") or 0)
    if inserted > 0 and failed <= 0 and queued <= 0:
        status = "success"
    elif inserted > 0 or duplicates > 0:
        status = "partial_success"
    else:
        status = "failed"
    message = f"写入 {inserted} 条，待审查 {queued} 条，去重跳过 {duplicates} 条"
    return {"name": source_name, "status": status, "message": message}


def import_daily_fault_summary_file(
    source_file,
    *,
    project_code,
    database_path=None,
    report_path=None,
    dry_run=False,
    timezone_default=TIMEZONE_DEFAULT,
):
    parsed = parse_daily_fault_summary(source_file)
    database_path = Path(database_path or get_db_path()).resolve()
    conn = create_db_connection(database_path, row_factory=True, enable_wal=True)
    ensure_ai_runtime_schema(conn)
    ai_service = DailyFaultSummaryAIService()
    batch_id = None
    resolved_report_path = Path(report_path).resolve() if report_path else None
    summary = {
        "inserted": 0,
        "duplicates_skipped": 0,
        "queue_items_created": 0,
        "station_proposals_created": 0,
        "rows_skipped": 0,
        "fail_count": 0,
    }
    row_results = []

    try:
        if not _supports_daily_fault_import(conn):
            raise RuntimeError("当前数据库未启用故障批次导入能力")

        project = get_project_row(conn, project_code)
        if not project:
            raise DailyFaultSummaryParseError(f"未找到项目: {project_code}")

        if not dry_run:
            batch_id = create_import_batch(
                conn,
                project_id=project["id"],
                source_type=SOURCE_TYPE,
                mode=IMPORT_MODE,
                file_count=1,
                report_path=None,
                timezone_default_used=timezone_default,
            )
            if resolved_report_path is None:
                resolved_report_path = database_path.parent / "import_reports" / f"{SOURCE_TYPE}_batch_{batch_id}.json"
            conn.execute(
                "UPDATE import_batches SET report_path = ? WHERE id = ?",
                (str(resolved_report_path), batch_id),
            )

        for entry in parsed["entries"]:
            description = build_fault_description(entry["section"], entry["problem_description"])
            fault_type = infer_fault_type(entry["section"], entry["problem_description"])
            ai_result = ai_service.analyze_entry(
                project_code=project_code,
                title=parsed["title"],
                source_date=parsed["source_date"],
                section=entry["section"],
                station_name=entry["station_name"],
                problem_description=entry["problem_description"],
            )
            if ai_result and ai_result.get("fault_type"):
                fault_type = ai_result["fault_type"]
            camera_location_text = _derive_camera_location_text(
                entry["problem_description"],
                ai_result,
            )
            match = resolve_station_match(conn, entry["station_name"], source_system=SOURCE_SYSTEM)
            source_record_key = build_source_record_key(
                project["code"],
                SOURCE_TYPE,
                raw_external_id=f"{parsed['source_date']}|{entry['station_name']}|{description}",
            )
            row_label = f"第{entry['row_index']}行 {entry['station_name']}"
            raw_context = {
                "title": parsed["title"],
                "section": entry["section"],
                "station_text": entry["station_name"],
                "station_token": entry["station_name"],
                "content": entry["problem_description"],
                "description": description,
                "fault_type": fault_type,
                "reporter_name": parsed["reporter_name"],
                "status": "open",
                "camera_location_text": camera_location_text,
                "parsed_time": parsed["source_date"],
                "raw_time": parsed["source_time_raw"],
                "source_timezone": timezone_default,
                "system_type": _system_type_for_project(project["code"]),
            }
            if ai_result:
                raw_context["ai_suggestion"] = ai_result
                raw_context["ai_confidence"] = ai_result.get("confidence")
                raw_context["ai_reason"] = ai_result.get("reason")

            if match["should_create_proposal"] and batch_id is not None and not dry_run:
                proposal_id = ensure_station_name_proposal(
                    conn,
                    import_batch_id=batch_id,
                    project_id=project["id"],
                    source_system=SOURCE_SYSTEM,
                    external_name=entry["station_name"],
                    candidate_station_id=match["proposal_candidate_station_id"],
                    confidence_score=match["confidence_score"],
                    raw_context=raw_context,
                )
                if proposal_id:
                    summary["station_proposals_created"] += 1

            if not match["matched"]:
                summary["fail_count"] += 1
                row_results.append(
                    {
                        "id": row_label,
                        "issue_type": "station_not_resolved",
                        "issue_detail": f"未识别站点：{entry['station_name']}",
                    }
                )
                if batch_id is not None and not dry_run:
                    queue_id = enqueue_fault_review_item(
                        conn,
                        import_batch_id=batch_id,
                        project_id=project["id"],
                        source_type=SOURCE_TYPE,
                        source_record_key_candidate=source_record_key,
                        raw_payload=raw_context,
                        issue_type="station_not_resolved",
                        issue_detail=f"未识别站点：{entry['station_name']}",
                        ai_suggestion=ai_result,
                        ai_confidence=ai_result.get("confidence") if ai_result else None,
                        ai_reason=ai_result.get("reason") if ai_result else None,
                    )
                    if queue_id:
                        summary["queue_items_created"] += 1
                continue

            if source_record_key and fault_source_key_exists(conn, source_record_key):
                summary["duplicates_skipped"] += 1
                row_results.append(
                    {
                        "id": row_label,
                        "status": "duplicate_skip",
                        "message": "已按来源键跳过重复记录",
                    }
                )
                continue

            semantic_duplicate_fault_id = find_semantic_duplicate_fault(
                conn,
                station_id=match["station_id"],
                project_id=project["id"],
                source_date=parsed["source_date"],
                fault_type=fault_type,
                description=description,
                camera_location_text=camera_location_text,
            )
            if semantic_duplicate_fault_id:
                summary["duplicates_skipped"] += 1
                row_results.append(
                    {
                        "id": row_label,
                        "status": "duplicate_skip",
                        "message": f"已按语义指纹跳过重复记录（故障ID {semantic_duplicate_fault_id}）",
                    }
                )
                continue

            camera_binding = resolve_camera_binding(
                conn,
                station_id=match["station_id"],
                project_id=project["id"],
                description=camera_location_text or entry["problem_description"],
            )
            if camera_binding:
                raw_context["camera_id"] = camera_binding["camera_id"]
                raw_context["camera_slot_id"] = camera_binding["slot_id"]
                raw_context["location"] = camera_binding["location_desc"]

            insert_fault_report(
                conn,
                station_id=match["station_id"],
                project_id=project["id"],
                batch_id=batch_id,
                project_code=project["code"],
                source_date=parsed["source_date"],
                source_time_raw=parsed["source_time_raw"],
                reporter_name=parsed["reporter_name"],
                description=description,
                fault_type=fault_type,
                source_record_key=source_record_key,
                camera_id=camera_binding["camera_id"] if camera_binding else None,
                camera_slot_id=camera_binding["slot_id"] if camera_binding else None,
                camera_location_text=raw_context.get("location") or camera_location_text,
                ai_confidence=ai_result.get("confidence") if ai_result else None,
                ai_trace_json=build_ai_trace(
                    ai_result,
                    provider=ai_service.provider_name,
                    model=ai_service.settings.model,
                    enabled=ai_service.enabled,
                ),
                timezone_default=timezone_default,
            )
            summary["inserted"] += 1
            row_results.append(
                {
                    "id": row_label,
                    "status": "inserted",
                    "message": f"{fault_type} / {description}",
                }
            )

        if not dry_run and batch_id is not None:
            update_import_batch_stats(
                conn,
                batch_id,
                success_count=summary["inserted"],
                fail_count=summary["fail_count"],
            )
            conn.commit()
        elif dry_run:
            conn.rollback()

        report = {
            "batch_id": batch_id,
            "project": project["code"],
            "source": str(parsed["source_path"]),
            "source_type": SOURCE_TYPE,
            "mode": IMPORT_MODE,
            "dry_run": dry_run,
            "source_date": parsed["source_date"],
            "timezone_default": timezone_default,
            "summary": summary,
            "files": [_build_file_result(parsed["source_path"].name, summary)],
            "rows": row_results,
            "aborted": False,
        }
        if resolved_report_path:
            resolved_report_path.parent.mkdir(parents=True, exist_ok=True)
            resolved_report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        return report
    except Exception:
        if not dry_run:
            conn.rollback()
        raise
    finally:
        conn.close()


def parse_args():
    parser = argparse.ArgumentParser(description="Import daily fault summary Excel into fault_reports")
    parser.add_argument("--source", required=True, help="Path to daily summary xls/xlsx file")
    parser.add_argument("--project", required=True, help="Target project code")
    parser.add_argument("--database", default=get_db_path(), help="Path to SQLite database")
    parser.add_argument("--report", help="Optional JSON report output path")
    parser.add_argument("--dry-run", action="store_true", help="Preview only, do not write database")
    parser.add_argument("--timezone-default", default=TIMEZONE_DEFAULT, help="Timezone label stored on imported rows")
    return parser.parse_args()


def main():
    args = parse_args()
    report = import_daily_fault_summary_file(
        args.source,
        project_code=args.project,
        database_path=args.database,
        report_path=args.report,
        dry_run=args.dry_run,
        timezone_default=args.timezone_default,
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
