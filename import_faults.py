#!/usr/bin/env python3
"""Generic historical fault importer for multi-project deployments."""

from __future__ import annotations

import argparse
import csv
import json
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from zoneinfo import ZoneInfo

import openpyxl

from import_review_support import (
    build_source_record_key,
    create_import_batch,
    enqueue_fault_review_item,
    ensure_station_name_proposal,
    fault_source_key_exists,
    get_columns,
    get_project_row,
    multi_project_import_enabled,
    resolve_station_match,
    update_import_batch_stats,
)
from utils import backup_sqlite_database, create_db_connection


DEFAULT_SOURCE_TYPE = "import_excel"
DEFAULT_TIMEZONE = "Asia/Shanghai"
FAIL_ON_ANY_FAILURE = "any_failure"
FAIL_ON_DUPLICATE = "duplicate"
FAIL_ON_ROWS_SKIPPED = "rows_skipped"

HEADER_ALIASES = {
    "external_id": {"externalid", "recordid", "sourceid", "id"},
    "station_name": {"station", "stationname", "stationtext", "substation"},
    "slot_code": {"slot", "slotcode", "cameraslotcode"},
    "project_device_code": {"devicecode", "projectdevicecode", "cameracode"},
    "camera_ip": {"cameraip", "ip", "ipaddress"},
    "fault_type_code": {"faulttypecode", "typecode"},
    "fault_type_label": {"faulttype", "faulttypelabel", "typelabel"},
    "description": {"description", "content", "detail", "message"},
    "occurred_at": {"occurredat", "faulttime", "createdat", "time"},
    "closed_at": {"closedat", "resolvedat", "finishedat"},
    "status": {"status", "faultstatus"},
    "handler_name": {"handler", "handlername", "owner"},
    "source_timezone": {"sourcetimezone", "timezone", "tz"},
}


@dataclass
class ImportStats:
    inserted: int = 0
    duplicates_skipped: int = 0
    queue_items_created: int = 0
    station_proposals_created: int = 0
    rows_skipped: int = 0
    fail_count: int = 0

    def to_dict(self) -> dict:
        return {
            "inserted": self.inserted,
            "duplicates_skipped": self.duplicates_skipped,
            "queue_items_created": self.queue_items_created,
            "station_proposals_created": self.station_proposals_created,
            "rows_skipped": self.rows_skipped,
            "fail_count": self.fail_count,
        }


class ImportAbortError(RuntimeError):
    def __init__(self, message: str, *, report: dict | None = None):
        super().__init__(message)
        self.report = report


def parse_fail_on_rules(value: str | None) -> set[str]:
    if not value:
        return set()
    rules = set()
    for part in str(value).split(","):
        normalized = part.strip().lower().replace("-", "_")
        if normalized:
            rules.add(normalized)
    return rules


def build_import_report(
    *,
    database: Path,
    source: Path,
    project_code: str,
    source_type: str,
    mode: str,
    dry_run: bool,
    stats: ImportStats,
    report_rows: list[dict],
    fail_on_rules: set[str],
    aborted: bool = False,
    error: str | None = None,
) -> dict:
    report = {
        "database": str(database),
        "source": str(source),
        "project": project_code,
        "source_type": source_type,
        "mode": mode,
        "dry_run": dry_run,
        "fail_on": sorted(fail_on_rules),
        "aborted": aborted,
        **stats.to_dict(),
        "rows": report_rows,
    }
    if error:
        report["error"] = error
    return report


def write_report_file(report_path: Path | None, report: dict) -> None:
    if not report_path:
        return
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def should_abort_import(*, fail_on_rules: set[str], reason: str, action: str) -> bool:
    if not fail_on_rules:
        return False
    normalized_reason = reason.strip().lower().replace("-", "_")
    normalized_action = action.strip().lower().replace("-", "_")
    if FAIL_ON_ANY_FAILURE in fail_on_rules and normalized_action in {"queue", "skip"}:
        return True
    if normalized_reason in fail_on_rules or normalized_action in fail_on_rules:
        return True
    if normalized_action == "duplicate_skip" and FAIL_ON_DUPLICATE in fail_on_rules:
        return True
    if normalized_action == "skip" and FAIL_ON_ROWS_SKIPPED in fail_on_rules:
        return True
    return False


def abort_import_if_needed(
    *,
    fail_on_rules: set[str],
    row_result: dict,
    reason: str,
    action: str,
) -> None:
    if not should_abort_import(fail_on_rules=fail_on_rules, reason=reason, action=action):
        return
    row_index = row_result.get("row_index")
    raise ImportAbortError(
        f"fail-on rule triggered: {reason} at row {row_index}" if row_index else f"fail-on rule triggered: {reason}"
    )


def normalize_header(value: str | None) -> str:
    return "".join(ch for ch in str(value or "").strip().lower() if ch.isalnum())


def map_headers(headers: Iterable[str | None]) -> dict[int, str]:
    mapping: dict[int, str] = {}
    for idx, header in enumerate(headers):
        normalized = normalize_header(header)
        if not normalized:
            continue
        for canonical, aliases in HEADER_ALIASES.items():
            if normalized == canonical or normalized in aliases:
                mapping[idx] = canonical
                break
    return mapping


def load_rows(source: Path) -> list[dict]:
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        workbook = openpyxl.load_workbook(source, data_only=True)
        try:
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if not rows:
            return []
        header_map = map_headers(rows[0])
        result = []
        for row_index, row in enumerate(rows[1:], start=2):
            payload = {"_row_index": row_index}
            for idx, canonical in header_map.items():
                payload[canonical] = row[idx] if idx < len(row) else None
            result.append(payload)
        return result

    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            rows = list(reader)
        if not rows:
            return []
        header_map = map_headers(rows[0])
        result = []
        for row_index, row in enumerate(rows[1:], start=2):
            payload = {"_row_index": row_index}
            for idx, canonical in header_map.items():
                payload[canonical] = row[idx] if idx < len(row) else None
            result.append(payload)
        return result

    raise ValueError(f"unsupported source file: {source}")


def load_type_mapping(mapping_path: Path | None) -> dict[str, str]:
    if not mapping_path:
        return {}
    with mapping_path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        mapping = {}
        for row in reader:
            source_value = (row.get("source_value") or row.get("source_label") or row.get("source") or "").strip()
            target_code = (row.get("target_code") or row.get("fault_type_code") or "").strip()
            if source_value and target_code:
                mapping[source_value] = target_code
        return mapping


def parse_source_timestamp(value, default_timezone: str) -> tuple[str | None, str | None]:
    if value in (None, ""):
        return None, None
    raw_text = str(value).strip()
    tz = ZoneInfo(default_timezone)

    if hasattr(value, "tzinfo") or hasattr(value, "hour"):
        if isinstance(value, datetime):
            dt = value
        else:
            dt = datetime.combine(value, datetime.min.time())
    else:
        text = raw_text.replace("/", "-").replace("T", " ")
        try:
            dt = datetime.fromisoformat(text)
        except ValueError:
            formats = [
                "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M",
                "%Y-%m-%d",
            ]
            dt = None
            for fmt in formats:
                try:
                    dt = datetime.strptime(text, fmt)
                    break
                except ValueError:
                    continue
            if dt is None:
                return None, raw_text

    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=tz)
    dt_utc = dt.astimezone(timezone.utc)
    return dt_utc.strftime("%Y-%m-%dT%H:%M:%SZ"), raw_text


def normalize_status(raw_status, *, closed_at: str | None) -> str:
    text = str(raw_status or "").strip().lower()
    if text in {"closed", "resolved", "done", "finished", "已关闭", "已完成"}:
        return "closed"
    if text in {"handling", "processing", "处理中", "待确认"}:
        return "handling"
    if closed_at:
        return "closed"
    return "open"


def get_current_fault_type_catalog(conn, project_id: int) -> tuple[dict[str, dict], dict[str, dict]]:
    project_row = conn.execute(
        "SELECT fault_type_version_id FROM projects WHERE id = ?",
        (project_id,),
    ).fetchone()
    version_id = project_row[0] if project_row else None
    if not version_id:
        return {}, {}

    rows = conn.execute(
        """
        SELECT id, version_id, type_code, type_label
        FROM project_fault_types
        WHERE version_id = ? AND is_active = 1
        """,
        (version_id,),
    ).fetchall()
    by_code = {}
    by_label = {}
    for row in rows:
        item = {
            "version_id": row[1] if not hasattr(row, "keys") else row["version_id"],
            "type_code": row[2] if not hasattr(row, "keys") else row["type_code"],
            "type_label": row[3] if not hasattr(row, "keys") else row["type_label"],
        }
        by_code[item["type_code"]] = item
        by_label[item["type_label"]] = item
    return by_code, by_label


def resolve_slot(conn, *, project_id: int, station_id: int, row: dict) -> tuple[int | None, str | None]:
    slot_code = str(row.get("slot_code") or "").strip()
    if slot_code:
        match = conn.execute(
            """
            SELECT id, slot_code
            FROM camera_slots
            WHERE station_id = ? AND project_id = ? AND slot_code = ?
            """,
            (station_id, project_id, slot_code),
        ).fetchone()
        if match:
            return (
                match[0] if not hasattr(match, "keys") else match["id"],
                match[1] if not hasattr(match, "keys") else match["slot_code"],
            )

    camera_columns = get_columns(conn, "cameras")
    project_device_code = str(row.get("project_device_code") or "").strip()
    if project_device_code and "project_camera_code" in camera_columns:
        rows = conn.execute(
            """
            SELECT DISTINCT c.slot_id
            FROM cameras c
            WHERE c.station_id = ?
              AND c.project_id = ?
              AND c.project_camera_code = ?
              AND c.slot_id IS NOT NULL
            """,
            (station_id, project_id, project_device_code),
        ).fetchall()
        if len(rows) == 1:
            return rows[0][0], None

    camera_ip = str(row.get("camera_ip") or "").strip()
    if camera_ip:
        rows = conn.execute(
            """
            SELECT DISTINCT c.slot_id
            FROM cameras c
            WHERE c.station_id = ?
              AND c.project_id = ?
              AND c.ip_address = ?
              AND c.slot_id IS NOT NULL
            """,
            (station_id, project_id, camera_ip),
        ).fetchall()
        if len(rows) == 1:
            return rows[0][0], None

    return None, slot_code or None


def resolve_fault_type(row: dict, *, by_code: dict, by_label: dict, type_mapping: dict) -> tuple[str | None, str | None, int | None]:
    raw_code = str(row.get("fault_type_code") or "").strip()
    raw_label = str(row.get("fault_type_label") or "").strip()

    if raw_code and raw_code in by_code:
        item = by_code[raw_code]
        return item["type_code"], item["type_label"], item["version_id"]

    if raw_label and raw_label in by_label:
        item = by_label[raw_label]
        return item["type_code"], item["type_label"], item["version_id"]

    mapped_code = None
    if raw_code and raw_code in type_mapping:
        mapped_code = type_mapping[raw_code]
    elif raw_label and raw_label in type_mapping:
        mapped_code = type_mapping[raw_label]

    if mapped_code and mapped_code in by_code:
        item = by_code[mapped_code]
        return item["type_code"], item["type_label"], item["version_id"]

    return None, raw_label or raw_code or None, None


def build_canonical_row(row: dict) -> str:
    canonical = {
        "station_name": str(row.get("station_name") or "").strip(),
        "slot_code": str(row.get("slot_code") or "").strip(),
        "project_device_code": str(row.get("project_device_code") or "").strip(),
        "fault_type_code": str(row.get("fault_type_code") or "").strip(),
        "fault_type_label": str(row.get("fault_type_label") or "").strip(),
        "description": str(row.get("description") or "").strip(),
        "occurred_at": str(row.get("occurred_at") or "").strip(),
    }
    return json.dumps(canonical, ensure_ascii=False, sort_keys=True)


def insert_fault_report(
    conn,
    *,
    project_row: dict,
    batch_id: int | None,
    slot_id: int,
    station_id: int,
    source_type: str,
    source_record_key: str,
    status: str,
    occurred_at_utc: str,
    closed_at_utc: str | None,
    raw_occurred_at: str | None,
    raw_timezone: str,
    fault_type_code: str | None,
    fault_type_label_snapshot: str | None,
    description: str,
    project_device_code: str | None,
    handler_name: str | None,
):
    columns = get_columns(conn, "fault_reports")
    insert_columns = [
        "station_id",
        "fault_type",
        "description",
        "reporter_name",
        "status",
        "created_at",
        "updated_at",
    ]
    values = [
        station_id,
        fault_type_label_snapshot or fault_type_code or "未命名故障",
        description,
        "历史导入",
        status,
        occurred_at_utc,
        closed_at_utc or occurred_at_utc,
    ]

    optional_fields = [
        ("closed_at", closed_at_utc),
        ("project_id", project_row["id"]),
        ("camera_slot_id", slot_id),
        ("fault_type_code", fault_type_code),
        ("fault_type_label_snapshot", fault_type_label_snapshot),
        ("fault_type_version_id", project_row.get("fault_type_version_id")),
        ("source_type", source_type),
        ("source_batch_id", str(batch_id) if batch_id is not None else None),
        ("source_record_key", source_record_key),
        ("project_device_code", project_device_code),
        ("source_time_raw", raw_occurred_at),
        ("source_timezone", raw_timezone),
        ("handler_name", handler_name),
    ]
    for column_name, value in optional_fields:
        if column_name in columns:
            insert_columns.append(column_name)
            values.append(value)

    placeholders = ", ".join(["?"] * len(insert_columns))
    conn.execute(
        f"""
        INSERT INTO fault_reports ({", ".join(insert_columns)})
        VALUES ({placeholders})
        """,
        values,
    )


def run_batch_import(
    *,
    database: Path,
    source: Path,
    project_code: str,
    source_type: str = DEFAULT_SOURCE_TYPE,
    mode: str = "best-effort",
    dry_run: bool = False,
    report_path: Path | None = None,
    type_mapping_path: Path | None = None,
    timezone_default: str = DEFAULT_TIMEZONE,
    fail_on: str = "",
) -> dict:
    conn = create_db_connection(database, row_factory=True, enable_wal=True)
    stats = ImportStats()
    report_rows = []
    fail_on_rules = parse_fail_on_rules(fail_on)

    try:
        if not multi_project_import_enabled(conn):
            raise RuntimeError("multi-project import schema is not enabled")

        project_row = get_project_row(conn, project_code)
        if not project_row:
            raise RuntimeError(f"project not found: {project_code}")

        project_meta = conn.execute(
            "SELECT fault_type_version_id FROM projects WHERE id = ?",
            (project_row["id"],),
        ).fetchone()
        project_row["fault_type_version_id"] = (
            project_meta[0] if project_meta else None
        )

        rows = load_rows(source)
        type_mapping = load_type_mapping(type_mapping_path)
        fault_types_by_code, fault_types_by_label = get_current_fault_type_catalog(conn, project_row["id"])

        if mode == "full-rollback" and not dry_run:
            conn.execute("BEGIN")

        batch_id = None
        if not dry_run:
            batch_id = create_import_batch(
                conn,
                project_id=project_row["id"],
                source_type=source_type,
                mode=mode,
                file_count=1,
                report_path=str(report_path) if report_path else None,
                timezone_default_used=timezone_default,
            )

        for row in rows:
            row_result = {"row_index": row.get("_row_index")}
            station_name = str(row.get("station_name") or "").strip()
            if not station_name:
                stats.rows_skipped += 1
                row_result["action"] = "skip"
                row_result["reason"] = "missing station_name"
                report_rows.append(row_result)
                abort_import_if_needed(
                    fail_on_rules=fail_on_rules,
                    row_result=row_result,
                    reason="missing_station_name",
                    action="skip",
                )
                continue

            occurred_at_utc, raw_occurred_at = parse_source_timestamp(row.get("occurred_at"), timezone_default)
            if not occurred_at_utc:
                stats.fail_count += 1
                row_result["action"] = "queue"
                row_result["reason"] = "invalid occurred_at"
                report_rows.append(row_result)
                if not dry_run and batch_id is not None:
                    enqueue_fault_review_item(
                        conn,
                        import_batch_id=batch_id,
                        project_id=project_row["id"],
                        source_type=source_type,
                        source_record_key_candidate=None,
                        raw_payload=row,
                        issue_type="invalid_timestamp",
                        issue_detail="Unable to parse occurred_at",
                    )
                    stats.queue_items_created += 1
                abort_import_if_needed(
                    fail_on_rules=fail_on_rules,
                    row_result=row_result,
                    reason="invalid_timestamp",
                    action="queue",
                )
                if mode == "full-rollback":
                    raise RuntimeError("full-rollback aborted by invalid timestamp")
                continue

            closed_at_utc, _ = parse_source_timestamp(row.get("closed_at"), timezone_default)
            source_timezone = str(row.get("source_timezone") or timezone_default).strip() or timezone_default
            status = normalize_status(row.get("status"), closed_at=closed_at_utc)

            external_id = str(row.get("external_id") or "").strip()
            canonical_row = build_canonical_row(row)
            if external_id:
                source_record_key = build_source_record_key(project_code, source_type, raw_external_id=external_id)
            else:
                source_record_key = None

            match = resolve_station_match(conn, station_name, source_system=source_type)
            if match["should_create_proposal"] and batch_id is not None and not dry_run:
                proposal_id = ensure_station_name_proposal(
                    conn,
                    import_batch_id=batch_id,
                    project_id=project_row["id"],
                    source_system=source_type,
                    external_name=station_name,
                    candidate_station_id=match["proposal_candidate_station_id"],
                    confidence_score=match["confidence_score"],
                    raw_context=row,
                )
                if proposal_id:
                    stats.station_proposals_created += 1

            if not match["matched"]:
                stats.fail_count += 1
                row_result["action"] = "queue"
                row_result["reason"] = "station_not_resolved"
                report_rows.append(row_result)
                if not dry_run and batch_id is not None:
                    enqueue_fault_review_item(
                        conn,
                        import_batch_id=batch_id,
                        project_id=project_row["id"],
                        source_type=source_type,
                        source_record_key_candidate=build_source_record_key(project_code, source_type, canonical_row=canonical_row),
                        raw_payload=row,
                        issue_type="station_not_resolved",
                        issue_detail=f"Unable to resolve station: {station_name}",
                    )
                    stats.queue_items_created += 1
                abort_import_if_needed(
                    fail_on_rules=fail_on_rules,
                    row_result=row_result,
                    reason="station_not_resolved",
                    action="queue",
                )
                if mode == "full-rollback":
                    raise RuntimeError("full-rollback aborted by station_not_resolved")
                continue

            slot_id, slot_code = resolve_slot(
                conn,
                project_id=project_row["id"],
                station_id=match["station_id"],
                row=row,
            )
            if not slot_id:
                stats.fail_count += 1
                row_result["action"] = "queue"
                row_result["reason"] = "slot_not_resolved"
                row_result["station_id"] = match["station_id"]
                report_rows.append(row_result)
                if not dry_run and batch_id is not None:
                    enqueue_fault_review_item(
                        conn,
                        import_batch_id=batch_id,
                        project_id=project_row["id"],
                        source_type=source_type,
                        source_record_key_candidate=source_record_key or build_source_record_key(project_code, source_type, canonical_row=canonical_row),
                        raw_payload={**row, "station_id": match["station_id"], "slot_code_candidate": slot_code},
                        issue_type="slot_not_resolved",
                        issue_detail="Unable to resolve slot by slot_code / device_code / camera_ip",
                    )
                    stats.queue_items_created += 1
                abort_import_if_needed(
                    fail_on_rules=fail_on_rules,
                    row_result=row_result,
                    reason="slot_not_resolved",
                    action="queue",
                )
                if mode == "full-rollback":
                    raise RuntimeError("full-rollback aborted by slot_not_resolved")
                continue

            fault_type_code, fault_type_label_snapshot, version_id = resolve_fault_type(
                row,
                by_code=fault_types_by_code,
                by_label=fault_types_by_label,
                type_mapping=type_mapping,
            )
            if not fault_type_label_snapshot:
                fault_type_label_snapshot = "未命名故障"

            if not fault_type_code:
                stats.fail_count += 1
                row_result["action"] = "queue"
                row_result["reason"] = "fault_type_not_mapped"
                report_rows.append(row_result)
                if not dry_run and batch_id is not None:
                    enqueue_fault_review_item(
                        conn,
                        import_batch_id=batch_id,
                        project_id=project_row["id"],
                        source_type=source_type,
                        source_record_key_candidate=source_record_key or build_source_record_key(project_code, source_type, canonical_row=canonical_row),
                        raw_payload={**row, "station_id": match["station_id"], "slot_id": slot_id},
                        issue_type="fault_type_not_mapped",
                        issue_detail="Unable to map fault type to current project catalog",
                    )
                    stats.queue_items_created += 1
                abort_import_if_needed(
                    fail_on_rules=fail_on_rules,
                    row_result=row_result,
                    reason="fault_type_not_mapped",
                    action="queue",
                )
                if mode == "full-rollback":
                    raise RuntimeError("full-rollback aborted by fault_type_not_mapped")
                continue

            project_row["fault_type_version_id"] = version_id
            candidate_key = source_record_key or build_source_record_key(project_code, source_type, canonical_row=canonical_row)
            if not external_id:
                stats.fail_count += 1
                row_result["action"] = "queue"
                row_result["reason"] = "source_record_key_unavailable"
                report_rows.append(row_result)
                if not dry_run and batch_id is not None:
                    enqueue_fault_review_item(
                        conn,
                        import_batch_id=batch_id,
                        project_id=project_row["id"],
                        source_type=source_type,
                        source_record_key_candidate=candidate_key,
                        raw_payload={**row, "station_id": match["station_id"], "slot_id": slot_id},
                        issue_type="source_record_key_unavailable",
                        issue_detail="Missing stable external_id; row requires manual review",
                    )
                    stats.queue_items_created += 1
                abort_import_if_needed(
                    fail_on_rules=fail_on_rules,
                    row_result=row_result,
                    reason="source_record_key_unavailable",
                    action="queue",
                )
                if mode == "full-rollback":
                    raise RuntimeError("full-rollback aborted by source_record_key_unavailable")
                continue

            if fault_source_key_exists(conn, source_record_key):
                stats.duplicates_skipped += 1
                row_result["action"] = "duplicate-skip"
                row_result["reason"] = "source_record_key_exists"
                report_rows.append(row_result)
                abort_import_if_needed(
                    fail_on_rules=fail_on_rules,
                    row_result=row_result,
                    reason="source_record_key_exists",
                    action="duplicate_skip",
                )
                continue

            if not dry_run:
                insert_fault_report(
                    conn,
                    project_row=project_row,
                    batch_id=batch_id,
                    slot_id=slot_id,
                    station_id=match["station_id"],
                    source_type=source_type,
                    source_record_key=source_record_key,
                    status=status,
                    occurred_at_utc=occurred_at_utc,
                    closed_at_utc=closed_at_utc,
                    raw_occurred_at=raw_occurred_at,
                    raw_timezone=source_timezone,
                    fault_type_code=fault_type_code,
                    fault_type_label_snapshot=fault_type_label_snapshot,
                    description=str(row.get("description") or "").strip() or fault_type_label_snapshot,
                    project_device_code=str(row.get("project_device_code") or "").strip() or None,
                    handler_name=str(row.get("handler_name") or "").strip() or None,
                )
            stats.inserted += 1
            row_result["action"] = "insert"
            row_result["station_id"] = match["station_id"]
            row_result["slot_id"] = slot_id
            row_result["fault_type_code"] = fault_type_code
            report_rows.append(row_result)

        if not dry_run:
            if batch_id is not None:
                update_import_batch_stats(
                    conn,
                    batch_id,
                    success_count=stats.inserted,
                    fail_count=stats.fail_count,
                )
            conn.commit()
    except Exception as exc:
        if not dry_run:
            conn.rollback()
        report = build_import_report(
            database=database,
            source=source,
            project_code=project_code,
            source_type=source_type,
            mode=mode,
            dry_run=dry_run,
            stats=stats,
            report_rows=report_rows,
            fail_on_rules=fail_on_rules,
            aborted=True,
            error=str(exc),
        )
        write_report_file(report_path, report)
        if isinstance(exc, ImportAbortError):
            exc.report = report
        raise
    finally:
        conn.close()

    report = build_import_report(
        database=database,
        source=source,
        project_code=project_code,
        source_type=source_type,
        mode=mode,
        dry_run=dry_run,
        stats=stats,
        report_rows=report_rows,
        fail_on_rules=fail_on_rules,
    )
    write_report_file(report_path, report)
    return report


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Import historical fault records")
    subparsers = parser.add_subparsers(dest="command", required=True)

    batch = subparsers.add_parser("batch", help="Import a spreadsheet or CSV in batch mode")
    batch.add_argument("--source", required=True, help="Path to xlsx/csv source file")
    batch.add_argument("--project", required=True, help="Target project code")
    batch.add_argument("--database", default="station_monitor.db", help="Path to SQLite database")
    batch.add_argument("--source-type", default=DEFAULT_SOURCE_TYPE, help="Source type tag")
    batch.add_argument("--mode", choices=["full-rollback", "best-effort"], default="best-effort")
    batch.add_argument("--type-mapping", help="Optional CSV mapping file")
    batch.add_argument("--timezone-default", default=DEFAULT_TIMEZONE, help="Default timezone for naive timestamps")
    batch.add_argument("--report", help="Optional JSON report output path")
    batch.add_argument("--fail-on", default="", help="Comma-separated fail-fast rules, e.g. station_not_resolved,any_failure")
    batch.add_argument("--dry-run", action="store_true", help="Preview only, do not write database")
    batch.add_argument("--skip-backup", action="store_true", help="Skip automatic SQLite backup before write import")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    database = Path(args.database).resolve()
    if not args.dry_run and not args.skip_backup:
        backup_path = backup_sqlite_database(database, label="import_faults")
        if backup_path:
            print(f"[备份] 已创建数据库备份: {backup_path}")
    try:
        report = run_batch_import(
            database=database,
            source=Path(args.source).resolve(),
            project_code=args.project,
            source_type=args.source_type,
            mode=args.mode,
            dry_run=args.dry_run,
            report_path=Path(args.report).resolve() if args.report else None,
            type_mapping_path=Path(args.type_mapping).resolve() if args.type_mapping else None,
            timezone_default=args.timezone_default,
            fail_on=args.fail_on,
        )
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    except ImportAbortError as exc:
        if exc.report:
            print(json.dumps(exc.report, ensure_ascii=False, indent=2))
        else:
            print(json.dumps({"error": str(exc)}, ensure_ascii=False, indent=2))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
